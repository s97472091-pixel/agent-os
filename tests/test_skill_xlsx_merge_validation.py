"""xlsx merge targets: an overlapping target is refused loudly, not written.

An intersecting merge target used to be written silently, leaving the workbook
with invalid intersecting merge ranges (Excel repairs such files on open)
while the run reported success; both scripts now refuse it with ``ValueError``,
matching the loud-failure contract of the malformed-range path, so the caller
gets an exact message and nothing is written. See issue #3280.
"""

from __future__ import annotations

import sys
from pathlib import Path
from typing import Any

import pytest
from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).resolve().parents[1]
SCRIPTS = ROOT / "src" / "agentos" / "skills" / "bundled" / "xlsx" / "scripts"


def _scripts() -> tuple[Any, Any]:
    sys.path.insert(0, str(SCRIPTS))
    try:
        import create_xlsx  # type: ignore[import-not-found]
        import edit_xlsx  # type: ignore[import-not-found]
    finally:
        sys.path.pop(0)
    return edit_xlsx, create_xlsx


def _merged_book(tmp_path: Path) -> Path:
    """A workbook with A1:B1 already merged."""
    wb = Workbook()
    ws = wb.worksheets[0]
    ws.title = "Form"
    ws["A1"] = "T"
    ws.merge_cells("A1:B1")
    path = tmp_path / "book.xlsx"
    wb.save(str(path))
    return path


def _ranges(ws: Any) -> list[str]:
    return sorted(str(r) for r in ws.merged_cells.ranges)


def test_overlapping_merge_is_refused(tmp_path: Path) -> None:
    edit_xlsx, _ = _scripts()
    wb = load_workbook(str(_merged_book(tmp_path)))

    with pytest.raises(ValueError, match="overlaps existing merge range"):
        edit_xlsx.apply_ops(wb, [{"op": "merge_cells", "sheet": "Form", "range": "B1:C1"}])

    assert _ranges(wb["Form"]) == ["A1:B1"]


def test_identical_merge_is_refused(tmp_path: Path) -> None:
    edit_xlsx, _ = _scripts()
    wb = load_workbook(str(_merged_book(tmp_path)))

    with pytest.raises(ValueError, match="overlaps existing merge range"):
        edit_xlsx.apply_ops(wb, [{"op": "merge_cells", "sheet": "Form", "range": "A1:B1"}])

    assert _ranges(wb["Form"]) == ["A1:B1"]


def test_malformed_range_still_fails_loudly(tmp_path: Path) -> None:
    edit_xlsx, _ = _scripts()
    wb = load_workbook(str(_merged_book(tmp_path)))

    with pytest.raises(ValueError, match="not a valid coordinate or range"):
        edit_xlsx.apply_ops(wb, [{"op": "merge_cells", "sheet": "Form", "range": "NOTARANGE"}])


def test_non_overlapping_merge_still_applies(tmp_path: Path) -> None:
    edit_xlsx, _ = _scripts()
    wb = load_workbook(str(_merged_book(tmp_path)))

    applied = edit_xlsx.apply_ops(wb, [{"op": "merge_cells", "sheet": "Form", "range": "D1:E1"}])

    assert applied == 1


def test_create_xlsx_rejects_overlapping_merges() -> None:
    _, create_xlsx = _scripts()

    with pytest.raises(ValueError, match="overlaps existing merge range"):
        create_xlsx.build(
            {
                "sheets": [
                    {
                        "name": "S",
                        "rows": [["a", "b", "c"]],
                        "merged": ["A1:B1", "B1:C1"],
                    }
                ]
            }
        )


def test_create_xlsx_still_accepts_non_overlapping_merges() -> None:
    _, create_xlsx = _scripts()

    wb = create_xlsx.build(
        {
            "sheets": [
                {
                    "name": "S",
                    "rows": [["a", "b", "c"]],
                    "merged": ["A1:B1", "D1:E1"],
                }
            ]
        }
    )

    assert _ranges(wb.worksheets[0]) == ["A1:B1", "D1:E1"]
