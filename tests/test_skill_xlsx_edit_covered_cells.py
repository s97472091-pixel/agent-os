"""xlsx edit_xlsx — a set_cell into a covered merged cell must not abort the batch.

A covered cell (inside a merged range, not its top-left anchor) is read-only in
openpyxl; assigning through it used to raise ``AttributeError`` out of
``apply_ops``, killing the whole run and losing every other op with it.
See issue #3277.
"""

from __future__ import annotations

import sys
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).resolve().parents[1]
SCRIPTS = ROOT / "src" / "agentos" / "skills" / "bundled" / "xlsx" / "scripts"


def _edit_xlsx() -> Any:
    sys.path.insert(0, str(SCRIPTS))
    try:
        import edit_xlsx  # type: ignore[import-not-found]
    finally:
        sys.path.pop(0)
    return edit_xlsx


def _make_book(tmp_path: Path) -> Path:
    """A workbook with a merged header row: A1:B1 merged, A1 = "Title"."""
    wb = Workbook()
    ws = wb.worksheets[0]
    ws.title = "Form"
    ws["A1"] = "Title"
    ws["A2"] = "a"
    ws["B2"] = "b"
    ws.merge_cells("A1:B1")
    path = tmp_path / "book.xlsx"
    wb.save(str(path))
    return path


def test_set_cell_into_covered_cell_is_skipped(tmp_path: Path) -> None:
    edit_xlsx = _edit_xlsx()
    path = _make_book(tmp_path)

    wb = load_workbook(str(path))
    applied = edit_xlsx.apply_ops(
        wb, [{"op": "set_cell", "sheet": "Form", "row": 1, "col": 2, "value": "X"}]
    )

    assert applied == 0
    out = tmp_path / "out.xlsx"
    wb.save(str(out))
    reloaded = load_workbook(str(out))
    assert reloaded["Form"]["A1"].value == "Title"
    assert reloaded["Form"]["B1"].value is None


def test_batch_survives_a_covered_cell(tmp_path: Path) -> None:
    edit_xlsx = _edit_xlsx()
    path = _make_book(tmp_path)

    wb = load_workbook(str(path))
    applied = edit_xlsx.apply_ops(
        wb,
        [
            {"op": "set_cell", "sheet": "Form", "row": 2, "col": 1, "value": "OK"},
            {"op": "set_cell", "sheet": "Form", "row": 1, "col": 2, "value": "X"},
            {"op": "set_cell", "sheet": "Form", "row": 2, "col": 2, "value": "Z"},
        ],
    )

    assert applied == 2
    out = tmp_path / "out.xlsx"
    wb.save(str(out))
    reloaded = load_workbook(str(out))
    assert reloaded["Form"]["A2"].value == "OK"
    assert reloaded["Form"]["B2"].value == "Z"


def test_set_cell_into_merge_anchor_still_applies(tmp_path: Path) -> None:
    edit_xlsx = _edit_xlsx()
    path = _make_book(tmp_path)

    wb = load_workbook(str(path))
    applied = edit_xlsx.apply_ops(
        wb, [{"op": "set_cell", "sheet": "Form", "row": 1, "col": 1, "value": "New Title"}]
    )

    assert applied == 1
    out = tmp_path / "out.xlsx"
    wb.save(str(out))
    reloaded = load_workbook(str(out))
    assert reloaded["Form"]["A1"].value == "New Title"
