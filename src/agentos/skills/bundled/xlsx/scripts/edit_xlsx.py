"""Apply cell edits to an existing `.xlsx`.

Operations:
    {"op": "set_cell", "sheet": "Q3", "row": 1, "col": 1, "value": "..."}
    {"op": "set_cell", "sheet": "Q3", "row": 2, "col": 2, "value": "=SUM(B3:B10)"}
    {"op": "set_cell", "sheet": "Q3", "row": 3, "col": 3, "value": "=hello", "as_text": true}
    {"op": "set_cell", "sheet": "Q3", "row": 4, "col": 1, "value": null}
    {"op": "rename_sheet", "old": "Sheet1", "new": "Summary"}
    {"op": "merge_cells", "sheet": "Q3", "range": "A1:C1"}

`value` semantics for `set_cell`:

* An explicit ``null`` **clears** the cell. It is the only way to express that
  in this op schema, and the cell's style is left alone.
* A **missing** ``value`` key is a malformed operation: it is skipped and not
  counted in ``applied``, so a typo cannot silently wipe data.
* ``0``, ``false`` and ``""`` are values, not absence, and are written as given.

``rename_sheet`` lands the sheet on exactly the name asked for, or does nothing.
A name another sheet already holds -- Excel compares sheet names without regard
to case -- is refused and not counted in ``applied``, because openpyxl would
otherwise store ``Summary1`` and report success.
"""

from __future__ import annotations

import argparse
import json
import sys
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.worksheet.cell_range import CellRange

# Bundled scripts run under AgentOS's own interpreter; the path insert only
# matters in a source checkout where the package is not installed (#2804).
_SRC_ROOT = str(Path(__file__).resolve().parents[5])
if _SRC_ROOT not in sys.path:
    sys.path.insert(0, _SRC_ROOT)
from agentos.skill_stdio import write_stdout as _write_stdout  # noqa: E402

# Distinguishes {"value": null} from an op with no "value" key at all.
# ``op.get("value")`` collapses both to None, which would make a malformed
# operation indistinguishable from a deliberate clear.
_MISSING = object()


def _coerce(value: Any, as_text: bool) -> Any:
    """Return the value to assign, honouring an explicit ``as_text`` request.

    ``as_text`` means "store exactly what I passed", so it suppresses the
    ISO-8601 coercion below as well as the formula interpretation. The cell
    *type* is what carries the distinction and that needs the cell object, so
    :func:`apply_ops` applies it after assignment; nothing is prepended to the
    data here. Excel's leading apostrophe is an input-mode escape rather than
    content, and writing it into the string left the cell holding ``'=hello``
    where the caller asked for ``=hello``.
    """
    if as_text:
        if isinstance(value, str) and value.startswith("'="):
            # ``SKILL.md`` offers ``'=hello`` and ``as_text: true`` as two
            # spellings of one request, so the two have to land on one cell.
            # Excel's leading apostrophe is the input escape for a
            # formula-looking value, so it is consumed here and carried as the
            # ``quotePrefix`` style flag by :func:`apply_ops` instead of being
            # stored as data. Scoped to ``'=``: a value that legitimately opens
            # with an apostrophe (``'tis``) keeps it.
            return value[1:]
        return value
    if isinstance(value, str) and len(value) >= 19 and value[10] == "T":
        try:
            return datetime.fromisoformat(value)
        except ValueError:
            return value
    return value


#: Every op kind ``apply_ops`` knows. An op outside this set is a caller
#: mistake, not a no-op: the ops file is written by the agent one step before
#: the call, so ``set-cell`` for ``set_cell`` is a routine slip.
OP_KINDS = ("set_cell", "rename_sheet", "merge_cells")


class OpsError(ValueError):
    """An ops file that cannot be used. Reported as ``error:`` / exit 2, never
    as a traceback: the caller passed bad input, the script did not break."""


def load_ops(path: Path) -> list[dict[str, Any]]:
    """Read and validate the ops file, or raise :class:`OpsError`.

    Validation happens before the workbook is opened, so an unusable ops file
    cannot leave a half-applied workbook behind, and ``--out`` is never touched.
    """
    try:
        raw = json.loads(path.read_text(encoding="utf-8"))
    except (UnicodeDecodeError, json.JSONDecodeError) as exc:
        raise OpsError(f"ops {path} is not valid JSON: {exc}") from exc
    if not isinstance(raw, list):
        raise OpsError(f"ops {path} must be a JSON array of operations, got {type(raw).__name__}")
    for index, op in enumerate(raw):
        if not isinstance(op, dict):
            raise OpsError(f"op {index} must be an object, got {type(op).__name__}")
        kind = op.get("op")
        if kind not in OP_KINDS:
            raise OpsError(
                f"op {index} has unknown kind {kind!r}; expected one of {', '.join(OP_KINDS)}"
            )
    return raw


def _free_temp_title(wb: Any) -> str:
    """A sheet title no sheet in *wb* currently holds, in any capitalisation."""
    taken = {name.casefold() for name in wb.sheetnames}
    index = 0
    while True:
        candidate = f"_rename_{index}"
        if candidate.casefold() not in taken:
            return candidate
        index += 1


def _rename_sheet(wb: Any, old: str, new: str) -> bool:
    """Rename *old* to exactly *new*; return whether that happened.

    openpyxl runs an assigned title through ``avoid_duplicate_name``, which
    compares case-insensitively against **every** sheet name -- the renamed
    sheet's own included -- and on a hit stores *new* with a number glued on
    rather than refusing. So renaming onto a name another sheet already held
    wrote ``Summary1``, and merely correcting a sheet's own capitalisation
    (``data`` -> ``Data``) wrote ``Data1``. Both reported ``applied``, and every
    later op addressing ``Summary`` then read and wrote the *other* sheet.

    A name another sheet holds is refused here, uncounted, the way this op list
    already treats a ``set_cell`` with no ``value``. A name only the renamed
    sheet itself holds is a legitimate request, so it goes through a free
    intermediate title: that clears the old spelling before the new one is
    claimed, leaving the title exactly as asked.
    """
    sheet = wb[old]
    if any(name != old and name.casefold() == new.casefold() for name in wb.sheetnames):
        return False
    if new == old:
        return True
    if new.casefold() == old.casefold():
        sheet.title = _free_temp_title(wb)
    sheet.title = new
    return True


def _ranges_intersect(a: CellRange, b: CellRange) -> bool:
    return not (
        a.max_row < b.min_row
        or a.min_row > b.max_row
        or a.max_col < b.min_col
        or a.min_col > b.max_col
    )


def _merge_or_reject(ws: Any, rng: str) -> None:
    """Merge *rng*, refusing a target that overlaps an existing merge.

    A malformed range already raises from :class:`CellRange` — the pinned
    loud-failure contract for a bad op — but an overlapping range used to be
    written silently, leaving the workbook with intersecting merge ranges
    (invalid content Excel repairs on open) while the run reported success.
    It is refused the same way so the caller gets an exact message to correct
    against and nothing is written.
    """
    target = CellRange(rng)
    for existing in ws.merged_cells.ranges:
        if _ranges_intersect(target, existing):
            raise ValueError(f"merge range {rng} overlaps existing merge range {existing}")
    ws.merge_cells(rng)


def apply_ops(wb: Any, ops: list[dict[str, Any]]) -> int:
    applied = 0
    for op in ops:
        if not isinstance(op, dict):
            continue
        kind = op.get("op")
        if kind == "set_cell":
            sheet_name = op.get("sheet")
            row = op.get("row")
            col = op.get("col")
            value = op.get("value", _MISSING)
            if sheet_name not in wb.sheetnames or row is None or col is None:
                continue
            if value is _MISSING:
                continue
            ws = wb[sheet_name]
            as_text = bool(op.get("as_text"))
            coerced = _coerce(value, as_text)
            # Assign through the property, not Worksheet.cell(value=...): that
            # helper ends with `if value is not None: cell.value = value`, so an
            # explicit null only *reads* the cell and the old value survives
            # while this loop still counts the edit as applied. Fetching the
            # cell first also leaves its style untouched.
            cell = ws.cell(row=int(row), column=int(col))
            cell.value = coerced
            if as_text and isinstance(coerced, str):
                # Assigning a string that starts with ``=`` makes openpyxl mark
                # the cell as a formula, so the string type has to be restored
                # afterwards. ``quotePrefix`` is the stored form of Excel's
                # apostrophe escape, which is why it belongs on the style and
                # not in the value.
                cell.data_type = "s"
                if coerced.startswith("="):
                    cell.quotePrefix = True
            applied += 1
        elif kind == "rename_sheet":
            old = op.get("old")
            new = op.get("new")
            if old in wb.sheetnames and isinstance(new, str) and _rename_sheet(wb, old, new):
                applied += 1
        elif kind == "merge_cells":
            sheet_name = op.get("sheet")
            rng = op.get("range")
            if sheet_name in wb.sheetnames and isinstance(rng, str):
                _merge_or_reject(wb[sheet_name], rng)
                applied += 1
    return applied


def _parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Edit an .xlsx via JSON op list.")
    parser.add_argument("input", type=Path)
    parser.add_argument("ops", type=Path)
    parser.add_argument("--out", type=Path, required=True)
    return parser.parse_args()


def main() -> int:
    args = _parse_args()
    if not args.input.is_file():
        print(f"error: input {args.input} not found", file=sys.stderr)
        return 2
    if not args.ops.is_file():
        print(f"error: ops {args.ops} not found", file=sys.stderr)
        return 2
    try:
        ops = load_ops(args.ops)
    except OpsError as exc:
        print(f"error: {exc}", file=sys.stderr)
        return 2
    wb = load_workbook(filename=str(args.input))
    applied = apply_ops(wb, ops)
    args.out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(args.out))
    _write_stdout(json.dumps({"applied": applied}, ensure_ascii=False) + "\n")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
